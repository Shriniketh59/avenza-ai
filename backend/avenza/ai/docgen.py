"""Generated documents: "create a document / PDF / Word file on X", or "convert this answer to PDF".

The answer is written by the normal pipeline (sources, citations, fact check), then saved as a .docx (python-docx)
or .pdf (fpdf2, DejaVu fonts so ₹ and accents print) with a Sources section, and offered as a download.
Files live under FILES_PATH/<user id>/ and are served only to their owner.
"""
import json
import re
import uuid
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path

from flask import Blueprint, current_app, g, jsonify, send_file

from ..tokens import require_auth

bp = Blueprint("files", __name__)

_VERB = r"(generate|create|make|write|prepare|draft|build|produce|give|get|send|i need|i want|export|convert|save|download|turn)"
_KIND = r"(word document|word file|word doc|docx|doc|document|pdf|pdf file|report file)"
_REQUEST = re.compile(rf"\b{_VERB}\b(?P<middle>.{{0,40}}?)\b{_KIND}s?\b", re.I)
# "convert this to word", "save it as MS Word": a bare "word" only counts after to/into/as/in
# (so "write a 500 word essay" is not a document request).
# "explain GST in pdf", "notes on SIP as a word file", "pdf on inflation", "docx about AI": the format named next
# to the topic, with or without a verb.
_FORMAT_NEAR = re.compile(
    r"\b(in|as|into)\s+(an?\s+)?(pdf|docx|word(\s+(file|format|document|doc))?)(\s+(file|format))?\b"
    r"|^\s*(a\s+)?(pdf|docx|word document|word file)\s+(on|about|for|of|regarding|explaining)\b",
    re.I,
)
# Questions about the formats themselves are not requests: "what is a pdf", "how do I convert pdf to word".
_ABOUT_FORMAT = re.compile(r"^\s*(what|how|why|which|is|are|can i|should i)\b.*\b(pdf|docx|word)\b.*\?$|^\s*(what is|what's) an? (pdf|docx)\b", re.I)
_TO_WORD = re.compile(rf"\b{_VERB}\b(?P<middle>.{{0,40}}?)\b(to|into|as|in)\s+(an?\s+)?(ms\s+)?word\b", re.I)
# Questions about an existing upload ("summarise my document"), unless they ask to export it ("…as a pdf").
# ("the pdf" alone is not an upload: "send me the pdf of this answer", "generate the docx on GST".)
_EXISTING = re.compile(r"\b(my|our|uploaded|attached|this|that)\s+(documents?|files?|pdfs?|docx)\b", re.I)
_EXPORT = re.compile(r"\b(as|into|to|in)\s+(a\s+)?(pdf|word|docx|doc|document)\b", re.I)
_CONVERT = re.compile(r"\b(convert|export|save|download|turn)\b", re.I)
# "export this answer", "download that": a file of the previous answer, Word by default.
_EXPORT_PREVIOUS = re.compile(r"^\s*(please\s+)?(export|download|save)\s+(this|that|it|the (above |last |previous )?(answer|response|reply))\b", re.I)
_LEAD = re.compile(r"^\s*(on|about|for|regarding|of|explaining|covering|that explains|which explains|describing|with)\b\s*", re.I)
_PRONOUNS = {"this", "that", "it", "above", "previous", "last", "the", "answer", "response", "reply", "explanation",
             "chat", "conversation", "as", "a", "an", "to", "into", "in", "format", "please", "pls", "me", "file",
             "and", "now", "can", "you", "for", "of", "my", "same", "word", "pdf", "doc", "docx", "document"}


@dataclass
class DocRequest:
    fmt: str  # "docx" | "pdf"
    topic: str  # what to write about ("" when converting the previous answer)
    from_previous: bool  # turn the last answer into a file, no new writing


def request(text: str) -> DocRequest | None:
    """The document this message asks for, or None for an ordinary message."""
    if _ABOUT_FORMAT.search(text):
        return None
    m = _REQUEST.search(text) or _TO_WORD.search(text) or _FORMAT_NEAR.search(text)
    fmt = "pdf" if re.search(r"\bpdf\b", text, re.I) else "docx"
    if not m:
        return DocRequest(fmt, "", True) if _EXPORT_PREVIOUS.search(text) else None
    if _EXISTING.search(text):
        if not _EXPORT.search(text) and not _CONVERT.search(text):
            return None  # "write a summary of my document": a question about an upload, answered in chat
        if not re.search(r"\b(this|that|it|above|previous|last)\b(?!\s+(document|file|pdf))", text, re.I):
            # "summarise my uploaded report as a pdf": answer about the upload, then save that answer
            return DocRequest(fmt, text.strip(), False)
    topic = (text[:m.start()] + " " + text[m.end():]).strip(" .?!")
    topic = re.sub(r"^\s*(can|could|would|will)\s+you\s+|^\s*(please|pls|kindly)\s+|\s+(please|pls)\s*$", "", topic, flags=re.I)
    topic = re.sub(r"\s+", " ", topic).strip()
    while _LEAD.match(topic):
        topic = _LEAD.sub("", topic, count=1)
    words = [w for w in re.findall(r"[a-z0-9]+", topic.lower()) if w not in _PRONOUNS]
    return DocRequest(fmt, topic.strip(), from_previous=not words)


TEMPLATE_RULE = (
    "Document: the user asked for a downloadable document. Write the complete document itself (it is saved as a "
    "file exactly as written): start with a '# ' title line, then a short introduction, '## ' section headings, "
    "bullets or tables where they help, and a conclusion. No chatty preamble or closing offer. Keep source "
    "citations like [1]. The file cannot show LaTeX: write formulas in plain text with symbols, like "
    "s = √( Σ(x − x̄)² / (n − 1) )."
)


# --- Markdown -> blocks ---------------------------------------------------------------------------------------

def _blocks(md: str) -> list[tuple]:
    """("h", level, text) | ("p", text) | ("ul", [items]) | ("ol", [items]) | ("table", rows) | ("code", text) | ("quote", text)."""
    out, para, lines, i = [], [], md.replace("\r", "").split("\n"), 0

    def flush():
        if para:
            out.append(("p", " ".join(para)))
            para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            flush()
            code, i = [], i + 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            out.append(("code", "\n".join(code)))
        elif m := re.match(r"^(#{1,6})\s+(.*)", s):
            flush()
            out.append(("h", len(m[1]), m[2].strip()))
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            out.append(("table", rows))
            continue
        elif m := re.match(r"^[-*+]\s+(.*)", s):
            flush()
            if out and out[-1][0] == "ul":
                out[-1][1].append(m[1])
            else:
                out.append(("ul", [m[1]]))
        elif m := re.match(r"^\d+[.)]\s+(.*)", s):
            flush()
            if out and out[-1][0] == "ol":
                out[-1][1].append(m[1])
            else:
                out.append(("ol", [m[1]]))
        elif s.startswith(">"):
            flush()
            out.append(("quote", s.lstrip("> ").strip()))
        elif not s:
            flush()
        else:
            para.append(s)
        i += 1
    flush()
    return out


_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|\[[^\]]+\]\([^)\s]+\))")


def _runs(text: str) -> list[tuple[str, str]]:
    """Split inline Markdown into (style, text): style is "", "b", "i" or "code"; links become "text (url)"."""
    out = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append(("b", part[2:-2]))
        elif part.startswith("`") and part.endswith("`"):
            out.append(("code", part[1:-1]))
        elif part.startswith("[") and "](" in part:
            label, url = part[1:].split("](", 1)
            out.append(("", f"{label} ({url[:-1]})"))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append(("i", part[1:-1]))
        else:
            out.append(("", part))
    return out


def _plain(text: str) -> str:
    return "".join(t for _s, t in _runs(text))


# --- Writers ---------------------------------------------------------------------------------------------------

def _docx(title: str, blocks: list[tuple], sources: list[dict]) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    def para(text: str, style: str | None = None):
        p = doc.add_paragraph(style=style)
        for kind, t in _runs(text):
            run = p.add_run(t)
            run.bold = kind == "b"
            run.italic = kind == "i"
            if kind == "code":
                run.font.name = "Consolas"
        return p

    doc.add_heading(title, 0)
    meta = doc.add_paragraph(f"Generated by AVENZA AI · {datetime.now():%d %B %Y}")
    meta.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    for b in blocks:
        if b[0] == "h":
            doc.add_heading(_plain(b[2]), min(b[1], 4))
        elif b[0] == "p":
            para(b[1])
        elif b[0] in ("ul", "ol"):
            for item in b[1]:
                para(item, "List Bullet" if b[0] == "ul" else "List Number")
        elif b[0] == "quote":
            para(b[1], "Intense Quote")
        elif b[0] == "code":
            p = doc.add_paragraph()
            run = p.add_run(b[1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif b[0] == "table" and b[1]:
            cols = max(len(r) for r in b[1])
            table = doc.add_table(rows=0, cols=cols)
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(b[1]):
                cells = table.add_row().cells
                for c in range(cols):
                    text = _plain(row[c]) if c < len(row) else ""
                    cells[c].text = text
                    if r == 0:
                        for run in cells[c].paragraphs[0].runs:
                            run.bold = True
    if sources:
        doc.add_heading("Sources", 1)
        for s in sources:
            para(s["line"])
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


_FONTS = Path(__file__).resolve().parent.parent / "assets" / "fonts"
# PDF fonts cannot draw emoji; drop them (and the invisible variation selector) instead of printing boxes.
_EMOJI = re.compile("[\U0001F000-\U0001FAFF☀-➿️‍]")


def _pdf(title: str, blocks: list[tuple], sources: list[dict]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_font("DejaVu", "", str(_FONTS / "DejaVuSans.ttf"))
    pdf.add_font("DejaVu", "B", str(_FONTS / "DejaVuSans-Bold.ttf"))
    pdf.add_font("DejaVu", "I", str(_FONTS / "DejaVuSans-Oblique.ttf"))
    pdf.add_font("DejaVuMono", "", str(_FONTS / "DejaVuSansMono.ttf"))
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin

    def text(t: str, size: float = 11, style: str = "", indent: float = 0, gap: float = 2.5):
        # fpdf2's markdown mode understands **bold** and __italic__.
        t = _EMOJI.sub("", t)
        t = "".join(f"**{x}**" if k == "b" else f"__{x}__" if k == "i" else x for k, x in _runs(t))
        pdf.set_font("DejaVu", style, size)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(width - indent, size * 0.5, t, markdown=True, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(gap)

    text(title, 20, "B", gap=1)
    pdf.set_text_color(107, 114, 128)
    text(f"Generated by AVENZA AI · {datetime.now():%d %B %Y}", 9, gap=5)
    pdf.set_text_color(17, 24, 39)
    sizes = {1: 16, 2: 14, 3: 12.5}
    for b in blocks:
        if b[0] == "h":
            pdf.ln(1.5)
            text(_plain(b[2]), sizes.get(b[1], 11.5), "B", gap=1.5)
        elif b[0] == "p":
            text(b[1])
        elif b[0] in ("ul", "ol"):
            for n, item in enumerate(b[1], 1):
                text(("•  " if b[0] == "ul" else f"{n}.  ") + item, indent=4, gap=1.2)
            pdf.ln(1.5)
        elif b[0] == "quote":
            pdf.set_text_color(75, 85, 99)
            text(b[1], 10.5, "I", indent=4)
            pdf.set_text_color(17, 24, 39)
        elif b[0] == "code":
            pdf.set_font("DejaVuMono", "", 9)
            pdf.set_fill_color(243, 244, 246)
            pdf.multi_cell(width, 4.6, _EMOJI.sub("", b[1]), fill=True, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2.5)
        elif b[0] == "table" and b[1]:
            cols = max(len(r) for r in b[1])
            pdf.set_font("DejaVu", "", 9.5)
            with pdf.table(width=width, line_height=5.2, first_row_as_headings=True) as table:
                for row in b[1]:
                    r = table.row()
                    for c in range(cols):
                        r.cell(_EMOJI.sub("", _plain(row[c])) if c < len(row) else "")
            pdf.ln(3)
    if sources:
        pdf.ln(2)
        text("Sources", 14, "B", gap=1.5)
        for s in sources:
            text(s["line"], 9.5, indent=2, gap=1)
    return bytes(pdf.output())


def _title(md: str, fallback: str) -> tuple[str, str]:
    """The '# Title' line (removed from the body), else a title made from the request."""
    m = re.search(r"^\s*#\s+(.+)$", md, re.M)
    if m:
        return _plain(m[1]).strip(), md[:m.start()] + md[m.end():]
    words = fallback.strip().rstrip("?.!") or "AVENZA AI document"
    return words[:1].upper() + words[1:80], md


def _dir(user_id: int) -> Path:
    path = Path(current_app.config["FILES_PATH"]) / str(user_id)
    path.mkdir(parents=True, exist_ok=True)
    return path


def build(md: str, fmt: str, user_id: int, topic: str, items: list[dict] = ()) -> dict:
    """Save md as a .docx/.pdf. Returns {"id", "name"}; the file is served at /api/files/<id>."""
    # Drop the fact-check banner and closing chatter lines a chat answer may carry.
    md = re.sub(r"\n+> ⚠️.*$", "", md.strip(), flags=re.S)
    title, body = _title(md, topic)
    cited = {int(n) for n in re.findall(r"\[(\d+)\]", body)}
    sources = []
    for i, it in enumerate(items, 1):
        if i in cited or not cited:
            label = it.get("title") or it.get("filename") or it.get("label") or "source"
            url = it.get("url") or ""
            sources.append({"line": f"[{i}] {label}" + (f" — {url}" if url else "")})
    blocks = _blocks(body)
    data = _pdf(title, blocks, sources) if fmt == "pdf" else _docx(title, blocks, sources)
    file_id = uuid.uuid4().hex
    slug = re.sub(r"[^A-Za-z0-9]+", "-", title).strip("-")
    if len(slug) > 60:
        slug = slug[:60].rsplit("-", 1)[0]  # whole words only: "...-GST-in-India", not "...-in-Indi"
    slug = slug or "AVENZA-document"
    name = f"{slug}.{fmt}"
    folder = _dir(user_id)
    (folder / f"{file_id}.{fmt}").write_bytes(data)
    (folder / f"{file_id}.json").write_text(json.dumps({"name": name, "created": datetime.now().isoformat()}))
    return {"id": file_id, "name": name, "size": len(data)}


_ID = re.compile(r"^[a-f0-9]{32}$")
_MIME = {"pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


@bp.get("/files/<file_id>")
@require_auth
def get_file(file_id: str):
    if not _ID.match(file_id):
        return jsonify(error="Not found."), 404
    folder = _dir(g.user["id"])
    found = [p for p in folder.glob(f"{file_id}.*") if p.suffix in (".pdf", ".docx")]
    if not found:
        return jsonify(error="Not found."), 404
    meta = json.loads((folder / f"{file_id}.json").read_text()) if (folder / f"{file_id}.json").exists() else {}
    path = found[0]
    return send_file(path, mimetype=_MIME[path.suffix[1:]], as_attachment=True,
                     download_name=meta.get("name") or path.name)
